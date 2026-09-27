"""Tests for the construction schedule (scheduling/): durations, CPM logic, calendar, exports, persistence."""
from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from ai.extraction import default_building_params
from detailed_mto import Options, build_project, compute
from engineering import plot_templates
from knowledge import load_knowledge_base
from models.schemas import ProjectInputs
from scheduling import ScheduleSettings, build_schedule, build_schedule_workbook, material_delivery_plan
from scheduling.engine import WorkCalendar
from scheduling.export import schedule_csv


def _params(pi):
    return plot_templates.build_template_params(pi) or default_building_params(
        wall_thickness_mm=pi.wall_thickness_mm, wall_material=pi.wall_material, unit_system=pi.unit_system)


@pytest.fixture(scope="module")
def kb():
    return load_knowledge_base()


def _result(kb, marla=5, scope=None):
    pi = ProjectInputs(project_name="Sched test", plot_marla=marla)
    opts = Options(scope=scope) if scope else None
    p = build_project(pi, _params(pi), kb, options=opts)
    return compute(p, kb, scope=scope)


@pytest.fixture(scope="module")
def res(kb):
    return _result(kb)


START = "2026-10-05"  # a Monday


@pytest.fixture(scope="module")
def sched(res):
    return build_schedule(res, ScheduleSettings(start_date=START))


# ------------------------------------------------------------------ calendar
def test_calendar_skips_sundays_and_holidays():
    cal = WorkCalendar(date(2026, 10, 5), [0, 1, 2, 3, 4, 5], ["2026-10-07"])
    assert all(d.weekday() != 6 for d in cal.dates[:60])
    assert date(2026, 10, 7) not in cal.dates
    assert cal.first_on_or_after(date(2026, 10, 11)) == cal.dates.index(date(2026, 10, 12))  # Sunday -> Monday
    assert cal.dates[cal.last_on_or_before(date(2026, 10, 11))] == date(2026, 10, 10)  # Sunday -> Saturday


# ------------------------------------------------------------------ network & CPM
def test_schedule_starts_on_start_date_and_is_realistic(sched):
    assert sched.start == date(2026, 10, 5)
    assert sched.activities[0].id == "PRE"
    assert sched.activities[-1].id == "HND" or max(sched.activities, key=lambda a: a.ef).id == "HND"
    # a 5 marla double-storey house: roughly 5-12 months with the default norms
    assert 150 <= sched.calendar_days <= 365


def test_every_activity_respects_its_predecessors(sched):
    by = {a.id: a for a in sched.activities}
    for a in sched.activities:
        assert a.duration >= 1
        assert a.finish >= a.start
        for pid, lag in a.preds:
            p = by[pid]
            assert a.start >= p.finish + timedelta(days=1 + lag), (a.id, pid, lag)


def test_no_work_on_non_working_days(res):
    s = build_schedule(res, ScheduleSettings(start_date=START, holidays=["2026-10-06", "2026-12-25"]))
    for a in s.activities:
        assert a.start.weekday() != 6 and a.finish.weekday() != 6
        assert a.start not in (date(2026, 10, 6), date(2026, 12, 25))


def test_critical_path_is_continuous_from_start_to_finish(sched):
    crit = sched.critical_path()
    assert crit and all(a.total_float == 0 for a in crit)
    ids = {a.id for a in crit}
    assert "PRE" in ids and "HND" in ids
    assert any(a.id.startswith("SLB-") for a in crit)  # the slabs drive a house programme
    for a in sched.activities:
        assert a.total_float >= 0


def test_floors_are_built_bottom_up_with_curing(sched):
    slb = sched.by_id("SLB-GRO")
    col1 = sched.by_id("COL-FIR")
    mas0 = sched.by_id("MAS-GRO")
    assert col1.start >= slb.finish + timedelta(days=1 + sched.settings.column_curing_days)
    assert mas0.start >= slb.finish + timedelta(days=1 + sched.settings.deshuttering_days)
    assert sched.by_id("PTI").start >= sched.by_id("PLI-GRO").finish + timedelta(days=sched.settings.plaster_drying_days)


def test_quantities_are_conserved_across_activities(res, sched):
    """Every scheduled work item's quantity is split across activities without loss or double counting
    (electrical points appear twice on purpose: rough-in and final wiring)."""
    tot = {}
    for a in sched.activities:
        for c in a.components:
            tot[c.wi_id] = tot.get(c.wi_id, 0) + c.qty
    wi = {w.wi_id: w.qty for w in res.scoped_work_items()}
    for k, v in tot.items():
        mult = 2 if k in {"WI-EL-01", "WI-EL-02", "WI-EL-03", "WI-EL-04", "WI-EL-05", "WI-EL-06"} else 1
        assert v == pytest.approx(wi[k] * mult, rel=1e-6), k


def test_productivity_gangs_and_overrides_change_durations(res, sched):
    slow = build_schedule(res, ScheduleSettings(start_date=START, productivity_pct=50))
    assert slow.finish > sched.finish
    base = sched.by_id("MAS-GRO").duration
    two = build_schedule(res, ScheduleSettings(start_date=START, gang_overrides={"MAS-GRO": 2}))
    assert two.by_id("MAS-GRO").duration < base
    fixed = build_schedule(res, ScheduleSettings(start_date=START, duration_overrides={"MAS-GRO": 30}))
    assert fixed.by_id("MAS-GRO").duration == 30 and fixed.by_id("MAS-GRO").overridden
    fast = build_schedule(res, ScheduleSettings(start_date=START, rate_overrides={"MAS|WI-MS-02": 400}))
    assert fast.by_id("MAS-GRO").duration < base


def test_bigger_house_takes_longer(kb, sched):
    big = build_schedule(_result(kb, marla=10), ScheduleSettings(start_date=START))
    assert big.finish > sched.finish


def test_scope_limits_the_schedule(kb):
    s = build_schedule(_result(kb, scope="Structure"), ScheduleSettings(start_date=START))
    ids = {a.id for a in s.activities}
    assert "FTG" in ids and "SLB-GRO" in ids
    assert "PTI" not in ids and "SAN" not in ids  # paint and sanitary fittings are outside a structure-only scope
    # links are bridged over removed activities: still one connected network ending at handover
    assert s.by_id("HND") is not None and s.by_id("HND").preds


def test_one_gang_per_trade_sequences_floors(res):
    on = build_schedule(res, ScheduleSettings(start_date=START, one_gang_per_trade=True))
    assert on.by_id("PLI-FIR").start > on.by_id("PLI-GRO").finish


# ------------------------------------------------------------------ outputs
def test_delivery_plan_orders_before_need(res, sched):
    rows = material_delivery_plan(res, sched, lead_days=5)
    assert len(rows) >= 0.9 * len(res.purchase_list())
    for r in rows:
        assert r["Order by"] == r["Needed on site"] - timedelta(days=5)
    cement = [r for r in rows if r["Code"] == "CON-001"]
    assert cement and cement[0]["Needed on site"] <= sched.by_id("FTG").start


def test_workbook_and_csv(res, sched):
    data = build_schedule_workbook(sched, "Sched test", material_delivery_plan(res, sched))
    wb = load_workbook(BytesIO(data))
    assert {"Schedule", "Gantt", "Delivery_Plan"} <= set(wb.sheetnames)
    ws = wb["Schedule"]
    assert ws.cell(row=5, column=3).value == "Activity"
    assert ws.max_row - 5 == len(sched.activities)
    csv = schedule_csv(sched).decode()
    assert csv.count("\n") == len(sched.activities) + 1


# ------------------------------------------------------------------ UI edit handling (pure functions)
def test_table_edits_become_overrides(res):
    import pandas as pd

    from ui.schedule_views import apply_activity_edits, apply_rate_edits

    s = ScheduleSettings(start_date=START, duration_overrides={"EXC": 9})
    base = pd.DataFrame([{"ID": "EXC", "Gangs": 1.0, "Duration (wd)": 9}, {"ID": "PLB", "Gangs": 1.0, "Duration (wd)": 3}])
    edited = base.copy()
    edited.loc[0, "Gangs"] = 2.0  # new gangs -> duration recalculated (fixed duration dropped)
    edited.loc[1, "Duration (wd)"] = 25
    out = apply_activity_edits(s, base, edited)
    assert out.gang_overrides == {"EXC": 2.0} and out.duration_overrides == {"PLB": 25}
    sched = build_schedule(res, out)
    assert sched.by_id("PLB").duration == 25 and not sched.by_id("EXC").overridden

    rb = pd.DataFrame([{"Key": "MAS|WI-MS-02", "Output per gang-day": 160.0, "Default": 160.0},
                       {"Key": "PLI|WI-PL-01", "Output per gang-day": 300.0, "Default": 220.0}])
    re = rb.copy()
    re.loc[0, "Output per gang-day"] = 200.0
    re.loc[1, "Output per gang-day"] = 220.0  # back to default -> override removed
    out2 = apply_rate_edits(ScheduleSettings(rate_overrides={"PLI|WI-PL-01": 300.0}), rb, re)
    assert out2.rate_overrides == {"MAS|WI-MS-02": 200.0}
