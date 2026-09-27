"""Excel export of the construction schedule: activity table, a bar-chart Gantt grid and the delivery plan."""
from __future__ import annotations

from datetime import timedelta
from io import BytesIO
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from scheduling.engine import WEEKDAY_NAMES, Schedule

FONT = "Arial"
NAVY, TEAL, RED, GREY = "1F3864", "0D9488", "C0392B", "D9E1F2"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
PHASE_FILL = PatternFill("solid", fgColor=GREY)
BAR = PatternFill("solid", fgColor=TEAL)
BAR_CRIT = PatternFill("solid", fgColor=RED)
FLOAT_FILL = PatternFill("solid", fgColor="E5E7EB")
thin = Side(style="thin", color="D0D7E2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(wrap_text=True, vertical="top")
DATE_FMT = "dd-mmm-yy"


def _hdr(ws, row: int, headers: List[str], widths: Optional[List[float]] = None) -> None:
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = CENTER
        c.border = BORDER
        if widths and i <= len(widths):
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]


def _cell(ws, r, c, v, bold=False, fmt=None, fill=None, color=None, align=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name=FONT, bold=bold, color=color)
    x.border = BORDER
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    x.alignment = align or WRAP
    return x


def build_schedule_workbook(sched: Schedule, project_name: str = "Project", delivery_rows: Optional[List[dict]] = None) -> bytes:
    wb = Workbook()
    s = sched.settings
    acts = sched.activities
    idx = {a.id: i for i, a in enumerate(acts)}

    # ---------------- Schedule sheet
    ws = wb.active
    ws.title = "Schedule"
    ws["A1"] = f"{project_name} - Construction schedule"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
    days = ", ".join(WEEKDAY_NAMES[d] for d in sorted(s.work_weekdays))
    ws["A2"] = (f"Start {sched.start:%d %b %Y} · Finish {sched.finish:%d %b %Y} · {sched.working_days} working days "
                f"({sched.calendar_days} calendar days, ~{sched.calendar_days / 30.4:.1f} months) · Working days: {days} · "
                f"Productivity {s.productivity_pct:g}% · Holidays: {len(s.holidays)}") if acts else "No activities"
    ws["A2"].font = Font(name=FONT, italic=True, color="555555")
    ws["A3"] = ("Durations = quantity from the take-off ÷ (output per gang-day × gangs). Lags are calendar days "
                "(curing / de-shuttering). Red = critical path (zero float). Planning estimate - verify with the contractor.")
    ws["A3"].font = Font(name=FONT, size=9, color="777777")
    heads = ["#", "ID", "Activity", "Phase", "Floor", "Duration (wd)", "Start", "Finish", "Total float (wd)",
             "Critical", "Predecessors (lag cal. days)", "Gangs", "Work content", "Duration basis"]
    _hdr(ws, 5, heads, [5, 9, 46, 22, 14, 11, 12, 12, 11, 9, 26, 7, 60, 60])
    r = 6
    for i, a in enumerate(acts, 1):
        preds = ", ".join(f"{idx[p] + 1}" + (f"+{lag}d" if lag else "") for p, lag in a.preds if p in idx)
        red = RED if a.critical else None
        vals = [i, a.id, a.name, a.phase_name, a.floor, a.duration, a.start, a.finish, a.total_float,
                "Yes" if a.critical else "", preds, a.gangs, a.work_text(), a.basis_text()]
        for c, v in enumerate(vals, 1):
            _cell(ws, r, c, v, bold=(c == 3 and a.critical), color=red if c in (3, 10) else None,
                  fmt=DATE_FMT if c in (7, 8) else None)
        r += 1
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(heads))}{max(r - 1, 5)}"

    # ---------------- Gantt sheet (one column per week)
    g = wb.create_sheet("Gantt")
    g["A1"] = f"{project_name} - Gantt chart (weekly)"
    g["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
    g["A2"] = "Teal = activity, red = critical path, grey = total float. Columns are weeks starting Monday."
    g["A2"].font = Font(name=FONT, size=9, color="777777")
    if acts:
        w0 = sched.start - timedelta(days=sched.start.weekday())
        last = max(a.finish + timedelta(days=0) for a in acts)
        weeks = []
        d = w0
        while d <= last:
            weeks.append(d)
            d += timedelta(days=7)
        fixed = ["#", "Activity", "Dur.", "Start", "Finish"]
        _hdr(g, 4, fixed, [5, 58, 6, 11, 11])
        for j, wk in enumerate(weeks):
            col = len(fixed) + 1 + j
            c = g.cell(row=4, column=col, value=wk)
            c.number_format = "dd-mmm"
            c.font = Font(name=FONT, bold=True, color="FFFFFF", size=8)
            c.fill = HDR_FILL
            c.alignment = Alignment(text_rotation=90, horizontal="center")
            g.column_dimensions[get_column_letter(col)].width = 3.2
            if wk.day <= 7:  # month label above the first week of each month
                m = g.cell(row=3, column=col, value=wk.strftime("%b %y"))
                m.font = Font(name=FONT, bold=True, size=8, color=NAVY)
        g.row_dimensions[4].height = 48
        r = 5
        cur_phase = None
        # grouped by phase (WBS view); "#" is the activity number on the Schedule sheet
        for a in sorted(acts, key=lambda x: (x.phase_name, x.es, x.ef)):
            i = idx[a.id] + 1
            if a.phase_name != cur_phase:
                cur_phase = a.phase_name
                for c in range(1, len(fixed) + len(weeks) + 1):
                    g.cell(row=r, column=c).fill = PHASE_FILL
                pc = g.cell(row=r, column=2, value=cur_phase)
                pc.font = Font(name=FONT, bold=True, color=NAVY)
                r += 1
            _cell(g, r, 1, i)
            _cell(g, r, 2, a.name, bold=a.critical, color=RED if a.critical else None,
                  align=Alignment(vertical="center", wrap_text=False))
            _cell(g, r, 3, a.duration, align=CENTER)
            _cell(g, r, 4, a.start, fmt=DATE_FMT)
            _cell(g, r, 5, a.finish, fmt=DATE_FMT)
            late_finish = sched.calendar[a.lf] if a.lf < len(sched.calendar) else a.finish
            for j, wk in enumerate(weeks):
                we = wk + timedelta(days=6)
                cell = g.cell(row=r, column=len(fixed) + 1 + j)
                cell.border = Border(left=Side(style="hair", color="E5E7EB"))
                if a.start <= we and a.finish >= wk:
                    cell.fill = BAR_CRIT if a.critical else BAR
                elif a.total_float and a.finish < wk <= late_finish:
                    cell.fill = FLOAT_FILL
            r += 1
        g.freeze_panes = g.cell(row=5, column=len(fixed) + 1)

    # ---------------- Delivery plan
    if delivery_rows:
        d = wb.create_sheet("Delivery_Plan")
        d["A1"] = f"{project_name} - Material delivery plan"
        d["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
        d["A2"] = "When each material must be on site (start of the first activity that uses it) and the latest order date."
        d["A2"].font = Font(name=FONT, size=9, color="777777")
        cols = ["Order by", "Needed on site", "Material", "Quantity to buy", "Buy unit", "First used in", "Trade", "Code"]
        _hdr(d, 4, cols, [12, 13, 44, 13, 10, 46, 20, 10])
        for i, row in enumerate(delivery_rows, 5):
            for c, k in enumerate(cols, 1):
                _cell(d, i, c, row.get(k), fmt=DATE_FMT if k in ("Order by", "Needed on site") else
                      ("#,##0.##" if k == "Quantity to buy" else None))
        d.freeze_panes = "C5"
        d.auto_filter.ref = f"A4:H{4 + len(delivery_rows)}"

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def schedule_csv(sched: Schedule) -> bytes:
    import csv
    import io

    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["ID", "Activity", "Phase", "Floor", "Duration (working days)", "Start", "Finish", "Total float",
                "Critical", "Predecessors", "Work content"])
    for a in sched.activities:
        w.writerow([a.id, a.name, a.phase_name, a.floor, a.duration, a.start.isoformat(), a.finish.isoformat(),
                    a.total_float, "Yes" if a.critical else "", " ".join(f"{p}+{lag}" if lag else p for p, lag in a.preds),
                    a.work_text()])
    return buf.getvalue().encode("utf-8")
